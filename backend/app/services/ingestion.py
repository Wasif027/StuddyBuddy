"""Document ingestion: parse → chunk → embed → persist.

Accepts the file types businesses actually use — ``.pdf``, ``.docx``, ``.xlsx`` /
``.xls``, ``.pptx`` — plus raw pasted text. Spreadsheets are additionally stored
sheet-by-sheet as queryable tables (:class:`DocumentTable`) for text-to-SQL, and
decks slide-by-slide (:class:`DocumentSlide`). Ingestion is idempotent on a
SHA-256 of the extracted text.
"""

from __future__ import annotations

import hashlib
import io
import re
import time
import uuid
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import cache
from app.core.logging import get_logger
from app.core.telemetry import tracer
from app.models.orm import Chunk, Document, DocumentSlide, DocumentStatus, DocumentTable
from app.services.chunking import chunk_text
from app.services.embeddings import embed_texts
from app.services.parsing import ParsedSlide, ParsedTable, ParsedUpload

logger = get_logger(__name__)
_tracer = tracer(__name__)

# Business formats only. Plain-text / markdown go through the "paste" flow.
SUPPORTED_UPLOAD_TYPES = {".pdf", ".docx", ".xlsx", ".xls", ".pptx"}

_SLIDE_HEADING_RE = re.compile(r"^slide\s+(\d+)", re.IGNORECASE)
_PAGE_HEADING_RE = re.compile(r"^page\s+(\d+)", re.IGNORECASE)
_SHEET_HEADING_RE = re.compile(r"^sheet:\s*(.+)$", re.IGNORECASE)


class UnsupportedFileType(ValueError):
    pass


class ParseError(ValueError):
    """A file we accept but couldn't read — message is user-facing."""


def _sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ----------------------------------------------------------------- file parsing
def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
    except (PdfReadError, ValueError, OSError, KeyError) as exc:
        raise ParseError(
            "couldn't read this PDF — it may be corrupted or password-protected"
        ) from exc
    text = "\n\n".join(f"## Page {i + 1}\n\n{t}" for i, t in enumerate(pages) if t)
    if not text.strip():
        raise ParseError(
            "this PDF has no selectable text — it looks like a scan. Text extraction "
            "from scanned images (OCR) isn't supported yet."
        )
    return text


def _parse_docx(data: bytes) -> str:
    import docx
    from docx.document import Document as _Docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document: _Docx = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, OSError, ValueError) as exc:
        raise ParseError(
            "couldn't read this Word file — it may be corrupted or not a real .docx"
        ) from exc
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name.lower().startswith("heading"):
            parts.append(f"## {text}")
        else:
            parts.append(text)
    # tables (python-docx keeps these separate from paragraphs)
    from app.services.parsing import markdown_table

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) >= 2:
            parts.append("\n" + markdown_table(rows[0], rows[1:], max_rows=50))
    return "\n\n".join(parts)


def parse_upload(filename: str, data: bytes) -> ParsedUpload:
    """Parse an uploaded business file into text (+ tables / slides where applicable)."""
    name = (filename or "document").rsplit("/", 1)[-1]
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    title = name.rsplit(".", 1)[0] if "." in name else name

    if ext == ".pdf":
        return ParsedUpload(title=title, kind="pdf", text=_parse_pdf(data))
    if ext == ".docx":
        return ParsedUpload(title=title, kind="docx", text=_parse_docx(data))
    if ext in {".xlsx", ".xls"}:
        from app.services.xlsx_parse import parse_workbook

        try:
            tables, preview = parse_workbook(name, data)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("couldn't read this spreadsheet — it may be corrupted") from exc
        if not tables:
            raise ParseError("this spreadsheet appears to be empty")
        return ParsedUpload(title=title, kind="xlsx", text=preview, tables=tables)
    if ext == ".pptx":
        from app.services.pptx_parse import parse_deck

        try:
            slides, deck_md = parse_deck(data)
        except Exception as exc:
            raise ParseError("couldn't read this presentation — it may be corrupted") from exc
        if not any(s.bullets or s.tables or s.title for s in slides):
            raise ParseError(
                "this presentation has no extractable text — it may be entirely images"
            )
        return ParsedUpload(title=title, kind="pptx", text=deck_md, slides=slides)
    raise UnsupportedFileType(f"unsupported file type: {ext or 'unknown'}")


# ------------------------------------------------------------- chunk location
def _chunk_location(kind: str, heading: str | None, slides_by_idx: dict[int, ParsedSlide]) -> dict:
    if not heading:
        return {}
    h = heading.strip()
    if kind == "pdf":
        m = _PAGE_HEADING_RE.match(h)
        return {"kind": "page", "page": int(m.group(1))} if m else {}
    if kind == "pptx":
        m = _SLIDE_HEADING_RE.match(h)
        if not m:
            return {}
        n = int(m.group(1))
        s = slides_by_idx.get(n)
        meta = {"kind": "slide", "slide": n}
        if s:
            meta.update(
                title=s.title,
                dataScore=s.data_score,
                hasChart=s.has_chart,
                hasTable=s.has_table,
            )
        return meta
    if kind == "xlsx":
        m = _SHEET_HEADING_RE.match(h)
        return {"kind": "sheet", "sheet": m.group(1).strip()} if m else {"kind": "sheet"}
    return {}


# --------------------------------------------------------------------- persist
def ingest_content(
    db: Session,
    *,
    user_id: str,
    title: str,
    content: str,
    category: str | None = None,
    source_type: str = "document",
    metadata: dict[str, Any] | None = None,
    tables: list[ParsedTable] | None = None,
    slides: list[ParsedSlide] | None = None,
) -> tuple[Document, int, bool]:
    started = time.perf_counter()
    content = content.strip()
    digest = _sha256(content)
    slides_by_idx = {s.index: s for s in (slides or [])}

    with _tracer.start_as_current_span("ingestion.ingest_content") as span:
        span.set_attribute("title", title)
        span.set_attribute("chars", len(content))
        span.set_attribute("kind", source_type)

        existing = db.execute(
            select(Document).where(
                Document.content_sha256 == digest, Document.user_id == user_id
            )
        ).scalar_one_or_none()
        if existing:
            span.set_attribute("deduplicated", True)
            return existing, len(existing.chunks), True

        doc = Document(
            id=str(uuid.uuid4()),
            user_id=user_id,
            title=title,
            category=category,
            source_type=source_type,
            status=DocumentStatus.PROCESSING,
            content_sha256=digest,
            char_count=len(content),
            metadata_json=metadata or {},
        )
        db.add(doc)
        db.flush()

        try:
            pieces = chunk_text(content)
            embeddings = embed_texts([p.text for p in pieces])
            for piece, vector in zip(pieces, embeddings, strict=True):
                loc = _chunk_location(source_type, piece.heading, slides_by_idx)
                db.add(
                    Chunk(
                        id=str(uuid.uuid4()),
                        document_id=doc.id,
                        chunk_index=piece.index,
                        text=piece.text,
                        heading=piece.heading,
                        token_count=piece.token_count,
                        embedding=vector,
                        metadata_json={**(piece.metadata or {}), **loc},
                    )
                )

            for pos, t in enumerate(tables or []):
                db.add(
                    DocumentTable(
                        id=str(uuid.uuid4()),
                        document_id=doc.id,
                        sheet_name=t.sheet_name,
                        sql_name=t.sql_name,
                        position=pos,
                        columns_json=[
                            {"name": c.name, "sqlName": c.sql_name, "type": c.type}
                            for c in t.columns
                        ],
                        row_count=t.row_count,
                        truncated=t.truncated,
                        rows_json=t.rows,
                    )
                )

            for s in slides or []:
                db.add(
                    DocumentSlide(
                        id=str(uuid.uuid4()),
                        document_id=doc.id,
                        index=s.index,
                        title=s.title,
                        bullets_json=s.bullets,
                        tables_json=s.tables,
                        notes=s.notes,
                        has_chart=s.has_chart,
                        has_table=s.has_table,
                        data_score=s.data_score,
                    )
                )

            doc.status = DocumentStatus.READY
            db.commit()
        except Exception as exc:  # pragma: no cover - defensive
            db.rollback()
            doc = db.get(Document, doc.id)
            if doc:
                doc.status = DocumentStatus.FAILED
                doc.error = str(exc)[:1000]
                db.commit()
            logger.error("ingestion_failed", title=title, error=str(exc))
            raise

        cache.invalidate_prefix("ekdp:")
        span.set_attribute("chunks", len(pieces))
        span.set_attribute("elapsed_ms", (time.perf_counter() - started) * 1000)
        return doc, len(pieces), False
