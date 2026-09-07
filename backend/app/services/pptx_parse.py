"""Parse ``.pptx`` decks into per-slide structured extracts.

Each slide → title, bullet lines, tables (markdown), speaker notes, and a
``data_score`` in [0, 1] estimating how data-rich the slide is (charts, tables,
dense figures, superlative titles). Used to (a) build retrieval text and (b) let
the reference panel surface the slides that carry the important numbers.
"""

from __future__ import annotations

import io
import re

from app.core.logging import get_logger
from app.services.parsing import ParsedSlide, markdown_table

logger = get_logger(__name__)

_NUMISH = re.compile(r"[$£€]\s?\d|\d[\d,]*\.?\d*\s?%|\b\d{2,}\b")
_SUPERLATIVE = re.compile(
    r"\b(top|highest|lowest|best|worst|most|least|record|growth|decline|"
    r"revenue|profit|loss|margin|forecast|target|kpi|results?|q[1-4]|fy\d{2})\b",
    re.IGNORECASE,
)


def _table_to_md(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    return markdown_table(rows[0], rows[1:], max_rows=30)


def _data_score(title: str | None, bullets: list[str], has_chart: bool, has_table: bool) -> float:
    score = 0.0
    if has_chart:
        score += 0.45
    if has_table:
        score += 0.35
    body = " ".join(bullets)
    figures = len(_NUMISH.findall(body))
    score += min(figures * 0.06, 0.4)
    if title and _SUPERLATIVE.search(title):
        score += 0.12
    if title and _NUMISH.search(title):
        score += 0.08
    return round(min(score, 1.0), 3)


def parse_deck(data: bytes) -> tuple[list[ParsedSlide], str]:
    """Return ``(slides, deck_markdown)``."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    slides: list[ParsedSlide] = []
    blocks: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text.strip() or None
        except (AttributeError, ValueError):
            title = None

        bullets: list[str] = []
        tables: list[str] = []
        has_chart = False
        has_table = False

        for shape in slide.shapes:
            if shape == getattr(slide.shapes, "title", None):
                continue
            if getattr(shape, "has_chart", False):
                has_chart = True
                try:
                    chart = shape.chart
                    cats = [str(c) for c in chart.plots[0].categories][:12]
                    for series in chart.series:
                        vals = list(series.values)[:12]
                        pairs = ", ".join(
                            f"{c}={v}" for c, v in zip(cats, vals, strict=False) if v is not None
                        )
                        bullets.append(f"[chart] {series.name}: {pairs}" if pairs else f"[chart] {series.name}")
                except (AttributeError, ValueError, IndexError):
                    bullets.append("[chart]")
                continue
            if getattr(shape, "has_table", False):
                has_table = True
                md = _table_to_md(shape.table)
                if md:
                    tables.append(md)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                continue
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip() or para.text.strip()
                    if line and line != title:
                        bullets.append(line)

        notes = None
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip() or None
        except (AttributeError, ValueError):
            notes = None

        score = _data_score(title, bullets, has_chart, has_table)
        slides.append(
            ParsedSlide(
                index=idx,
                title=title,
                bullets=bullets,
                tables=tables,
                notes=notes,
                has_chart=has_chart,
                has_table=has_table,
                data_score=score,
            )
        )

        parts = [f"## Slide {idx}" + (f" — {title}" if title else "")]
        parts += [f"- {b}" for b in bullets]
        parts += tables
        if notes:
            parts.append(f"\n_Notes: {notes}_")
        blocks.append("\n".join(parts))

    return slides, "\n\n".join(blocks)
