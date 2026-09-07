"""Parse ``.xlsx`` / ``.xls`` workbooks into queryable :class:`ParsedTable`s.

Each worksheet becomes one table: a detected header row, inferred column types,
and JSON-safe rows (capped at ``settings.analysis_max_rows``). A compact markdown
preview per sheet is also produced so the sheet is still findable by ordinary
retrieval ("what columns are in the orders sheet?").
"""

from __future__ import annotations

import io
import re

from app.core.config import get_settings
from app.core.logging import get_logger
from app.services.parsing import (
    ParsedColumn,
    ParsedTable,
    infer_type,
    json_safe,
    markdown_table,
    sanitise_identifier,
)

logger = get_logger(__name__)
settings = get_settings()

_MAX_COLS = 100


def _header_index(rows: list[list]) -> int:
    """Pick the first mostly-populated row as the header."""
    for i, row in enumerate(rows[:15]):
        filled = [c for c in row if c not in (None, "")]
        if len(filled) >= max(2, len(row) * 0.5) and all(
            not isinstance(c, (int, float)) or isinstance(c, bool) for c in filled
        ):
            return i
    return 0


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (name or "sheet").strip().lower()).strip("_") or "sheet"


def _rows_from_xlsx(data: bytes) -> list[tuple[str, list[list]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[tuple[str, list[list]]] = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c not in (None, "") for c in r)]
            if rows:
                out.append((ws.title, rows))
    finally:
        wb.close()
    return out


def _rows_from_xls(data: bytes) -> list[tuple[str, list[list]]]:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, engine="xlrd")
    out: list[tuple[str, list[list]]] = []
    for name, df in sheets.items():
        rows = df.where(df.notna(), None).values.tolist()
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if rows:
            out.append((str(name), rows))
    return out


def parse_workbook(filename: str, data: bytes) -> tuple[list[ParsedTable], str]:
    """Return ``(tables, preview_markdown)``."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "xlsx"
    raw_sheets = _rows_from_xls(data) if ext == "xls" else _rows_from_xlsx(data)

    tables: list[ParsedTable] = []
    previews: list[str] = []
    taken_sql_names: set[str] = set()
    cap = settings.analysis_max_rows

    for sheet_name, rows in raw_sheets:
        hidx = _header_index(rows)
        header = rows[hidx]
        width = min(len(header), _MAX_COLS)
        data_rows_all = rows[hidx + 1 :]

        col_taken: set[str] = set()
        sheet_slug = _slug(sheet_name)
        sql_name = sanitise_identifier(sheet_name, taken_sql_names, fallback="sheet")

        # normalise every row to the header width, JSON-safe
        norm_rows: list[list] = []
        for r in data_rows_all:
            cells = [json_safe(c) for c in list(r)[:width]]
            cells += [None] * (width - len(cells))
            norm_rows.append(cells)

        truncated = len(norm_rows) > cap
        kept_rows = norm_rows[:cap]

        columns: list[ParsedColumn] = []
        for ci in range(width):
            raw_name = header[ci]
            name = str(raw_name).strip() if raw_name not in (None, "") else f"column_{ci + 1}"
            col_sql = sanitise_identifier(name, col_taken, fallback=f"{sheet_slug}_col")
            ctype = infer_type([r[ci] for r in kept_rows])
            columns.append(ParsedColumn(name=name, sql_name=col_sql, type=ctype))

        tables.append(
            ParsedTable(
                sheet_name=sheet_name,
                sql_name=sql_name,
                columns=columns,
                rows=kept_rows,
                row_count=len(norm_rows),
                truncated=truncated,
            )
        )

        previews.append(
            f"## Sheet: {sheet_name}\n\n"
            f"Columns: {', '.join(f'{c.name} ({c.type})' for c in columns)}\n"
            f"Rows: {len(norm_rows)}\n\n"
            + markdown_table([c.name for c in columns], kept_rows, max_rows=15)
        )

    preview_md = "\n\n---\n\n".join(previews) if previews else "(empty workbook)"
    return tables, preview_md
